#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx", "Pillow", "lxml", "latex2mathml", "mathml2omml"]
# ///

# ─── How to run ───
# python3 check_harness.py

"""Deterministic checks for the report style CLI. No external report is required."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from check_specs import figure_spec, manuscript, ref_spec, valid_spec
from check_style import assert_geometry, assert_style, prefix_preserved
from docx import Document
from docx.shared import Cm
from PIL import Image

SKILL = Path(__file__).resolve().parent.parent
CLI = SKILL / "scripts" / "build_report.py"
TEMPLATE = SKILL / "template" / "report-style.docx"
ALLOW = {
    "word/document.xml",
    "word/_rels/document.xml.rels",
    "[Content_Types].xml",
    "docProps/app.xml",
    "docProps/core.xml",
}
SENTINEL = "ORIGINAL_BODY_SENTINEL_9f3a"


def run_cli(args: list[str]) -> subprocess.CompletedProcess[str]:
    """Run the real CLI entry point."""
    return subprocess.run(
        [sys.executable, str(CLI), *args],
        check=False,
        capture_output=True,
        text=True,
        timeout=180,
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def part(path: Path, name: str) -> bytes:
    with zipfile.ZipFile(path) as archive:
        return archive.read(name)


def media_hashes(path: Path) -> dict[str, str]:
    found: dict[str, str] = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/") or Path(name).name.startswith("thumbnail."):
                found[name] = hashlib.sha256(archive.read(name)).hexdigest()
    return found


def write_png(path: Path, size: tuple[int, int], color: tuple[int, int, int]) -> None:
    Image.new("RGB", size, color).save(path)


def build(directory: Path, source: Path, mode: str, *, extra: list[str] | None = None) -> tuple[Path, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    wide = directory / "wide.png"
    tall = directory / "tall.png"
    write_png(wide, (2000, 500), (20, 40, 60))
    write_png(tall, (400, 2000), (60, 40, 20))
    spec = directory / "manuscript.json"
    spec.write_text(manuscript(wide, tall), encoding="utf-8")
    output = directory / f"{mode}.docx"
    before = sha256(source)
    args = ["--source", str(source), "--manuscript", str(spec), "--output", str(output), "--mode", mode]
    if extra:
        args.extend(extra)
    result = run_cli(args)
    assert result.returncode == 0, result.stderr
    assert sha256(source) == before
    audit = output.parent / f"{output.name}.audit.json"
    assert audit.is_file()
    return output, audit


def synthetic_source(directory: Path, images: list[tuple[str, tuple[int, int], tuple[int, int, int]]]) -> Path:
    dest = directory / "source.docx"
    shutil.copyfile(TEMPLATE, dest)
    document = Document(str(dest))
    document.add_paragraph(SENTINEL)
    for name, size, color in images:
        path = directory / name
        write_png(path, size, color)
        document.add_paragraph().add_run().add_picture(str(path), width=Cm(3))
    document.save(str(dest))
    return dest


def test_help_lists_modes() -> None:
    result = run_cli(["--help"])
    assert result.returncode == 0
    assert "--mode" in result.stdout
    assert "--source" in result.stdout
    assert "--manuscript" in result.stdout
    assert "--output" in result.stdout
    assert "--ledger" in result.stdout


def test_new_mode_preserves_style(tmp: Path) -> None:
    output, audit_path = build(tmp, TEMPLATE, "new")
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["mode"] == "new"
    assert audit["body_policy"] == "replaced"
    assert audit["styles_unchanged"] is True
    assert audit["theme_unchanged"] is True
    assert audit["section_geometry_unchanged"] is True
    assert audit["leaked_source_text_count"] == 0
    assert audit["disallowed_package_parts"] == []
    assert part(output, "word/styles.xml") == part(TEMPLATE, "word/styles.xml")
    assert part(output, "word/theme/theme1.xml") == part(TEMPLATE, "word/theme/theme1.xml")
    blob = part(output, "word/document.xml")
    assert_geometry(blob)
    assert_style(blob)
    placed = {sha256(tmp / "wide.png"), sha256(tmp / "tall.png")}
    assert set(media_hashes(output).values()) == placed


def test_new_mode_drops_unrequested_and_keeps_requested(tmp: Path) -> None:
    source = synthetic_source(
        tmp,
        [("kept.png", (40, 20), (10, 20, 30)), ("dropped.png", (30, 30), (200, 10, 10))],
    )
    before = sha256(source)
    with zipfile.ZipFile(source) as archive:
        names = [name for name in archive.namelist() if name.startswith("word/media/")]
        kept_blob = archive.read(names[0])
        dropped = hashlib.sha256(archive.read(names[1])).hexdigest()
    reused = tmp / "reused.png"
    reused.write_bytes(kept_blob)
    kept = hashlib.sha256(kept_blob).hexdigest()
    spec = tmp / "reuse.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "p", "text": "See Figure K."},
                    {
                        "type": "figure",
                        "id": "K",
                        "mention": "Figure K",
                        "asset": "kept",
                        "source": "this document",
                        "caption": "Figure K. this document.",
                    },
                ],
                "assets": {"kept": {"path": str(reused)}},
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "new.docx"
    result = run_cli(
        ["--source", str(source), "--manuscript", str(spec), "--output", str(output), "--mode", "new"]
    )
    assert result.returncode == 0, result.stderr
    blob = part(output, "word/document.xml")
    assert SENTINEL.encode() not in blob
    hashes = set(media_hashes(output).values())
    assert kept in hashes
    assert dropped not in hashes
    assert sha256(source) == before


def test_append_preserves_package(tmp: Path) -> None:
    source = synthetic_source(tmp, [("original.png", (24, 24), (1, 2, 3))])
    before_parts = media_hashes(source)
    output, audit_path = build(tmp / "append-run", source, "append")
    audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert audit["mode"] == "append"
    assert audit["original_body_xml_prefix_preserved"] is True
    assert audit["original_media_retained"] is True
    assert audit["styles_unchanged"] is True
    assert audit["removed_parts"] == []
    assert set(audit["changed_parts"]) <= ALLOW
    assert audit["disallowed_package_parts"] == []
    assert all(name.startswith("word/media/") for name in audit["added_parts"])
    assert prefix_preserved(part(source, "word/document.xml"), part(output, "word/document.xml"))
    assert SENTINEL.encode() in part(output, "word/document.xml")
    after_parts = media_hashes(output)
    for name, digest in before_parts.items():
        assert after_parts[name] == digest


def test_default_source_without_figures(tmp: Path) -> None:
    spec = tmp / "note.json"
    spec.write_text(
        json.dumps({"blocks": [{"type": "h", "level": 1, "text": "NOTE"}, {"type": "p", "text": "No figure and no citation."}]}),
        encoding="utf-8",
    )
    output = tmp / "note.docx"
    before = sha256(TEMPLATE)
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    assert result.returncode == 0, result.stderr
    assert sha256(TEMPLATE) == before
    assert part(output, "word/styles.xml") == part(TEMPLATE, "word/styles.xml")
    assert b"NOTE" in part(output, "word/document.xml")


def test_reject_malformed_missing_and_overwrite(tmp: Path) -> None:
    tiny = tmp / "tiny.docx"
    write_png(tmp / "ok.png", (8, 8), (1, 2, 3))
    Document().save(str(tiny))
    before = sha256(tiny)
    png_before = sha256(tmp / "ok.png")
    bad = tmp / "bad.json"
    bad.write_text("{", encoding="utf-8")
    malformed = run_cli(["--source", str(tiny), "--manuscript", str(bad), "--output", str(tmp / "out.docx"), "--mode", "new"])
    missing_figure = _reject(tmp, tiny, figure_spec(tmp, asset=str(tmp / "absent.png"), previous="Figure F is here.", caption="Figure F. this document."))
    missing_mention = _reject(tmp, tiny, figure_spec(tmp, asset=str(tmp / "ok.png"), previous="No token here.", caption="Figure F. this document."))
    missing_source = _reject(tmp, tiny, figure_spec(tmp, asset=str(tmp / "ok.png"), previous="Figure F is here.", caption="Figure F without provenance."))
    missing_ref = _reject(tmp, tiny, ref_spec(tmp, cite=False))
    unknown_ref = _reject(tmp, tiny, ref_spec(tmp, cite=True, extra="[MISSING]"))
    existing = tmp / "exists.docx"
    existing.write_bytes(b"keep")
    valid = valid_spec(tmp)
    overwrite = run_cli(["--source", str(tiny), "--manuscript", str(valid), "--output", str(existing), "--mode", "new"])
    collision = run_cli(["--source", str(tiny), "--manuscript", str(valid), "--output", str(tiny), "--mode", "append"])
    asset_hit = run_cli(["--source", str(tiny), "--manuscript", str(valid), "--output", str(tmp / "ok.png"), "--mode", "new"])
    assert malformed.returncode == 2 and "malformed_manuscript" in malformed.stderr
    assert "missing_figure" in missing_figure
    assert "missing_figure_mention" in missing_mention
    assert "missing_figure_source" in missing_source
    assert "missing_reference" in missing_ref
    assert "missing_reference" in unknown_ref
    assert overwrite.returncode == 2 and "output_exists" in overwrite.stderr
    assert existing.read_bytes() == b"keep"
    assert collision.returncode == 2 and "output_collides" in collision.stderr
    assert asset_hit.returncode == 2 and "output_collides" in asset_hit.stderr
    assert sha256(tiny) == before
    assert sha256(tmp / "ok.png") == png_before


def test_append_without_source_errors(tmp: Path) -> None:
    spec = tmp / "note.json"
    spec.write_text('{"blocks":[{"type":"p","text":"Note."}]}', encoding="utf-8")
    output = tmp / "out.docx"
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "append"])
    assert result.returncode == 2 and "missing_source" in result.stderr
    assert not output.exists()


def test_windows_path_is_preserved(tmp: Path) -> None:
    """Literal prose paths survive the actual CLI and Word text round trip."""
    # Given: ordinary technical prose, not an equation.
    text = r"The input is C:\data\input.csv."
    spec = tmp / "path.json"
    spec.write_text(json.dumps({"blocks": [{"type": "p", "text": text}]}), encoding="utf-8")
    output = tmp / "path.docx"
    # When: the real CLI builds it.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the machine-consumed manuscript value is preserved, not rewritten.
    assert result.returncode == 0, result.stderr
    assert Document(str(output)).paragraphs[0].text == text


def test_indexing_is_not_citation(tmp: Path) -> None:
    """Identifiers and chained indexing need no fabricated bibliography."""
    # Given: literal indexing without reference entries.
    text = "Read samples[0] and matrix[row][col] as indexed values."
    spec = tmp / "index.json"
    spec.write_text(json.dumps({"blocks": [{"type": "p", "text": text}]}), encoding="utf-8")
    output = tmp / "index.docx"
    # When: the real CLI builds it.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the original literal expression survives.
    assert result.returncode == 0, result.stderr
    assert Document(str(output)).paragraphs[0].text == text


def test_latex_equation_is_rejected(tmp: Path) -> None:
    """Equation lines still reject raw LaTeX. The conversion source is the latex field."""
    # Given: unsupported equation syntax in the visible line.
    spec = tmp / "latex.json"
    spec.write_text(
        json.dumps({"blocks": [{"type": "eq", "lines": [r"\frac{1}{2}"], "source": "this document"}]}),
        encoding="utf-8",
    )
    output = tmp / "latex.docx"
    # When: the real CLI attempts the build.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: raw LaTeX is not written as an equation.
    assert result.returncode == 2 and "latex_marker" in result.stderr
    assert not output.exists()


MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def test_linear_equation_preserves_source(tmp: Path) -> None:
    """A line without latex stays an editable equation and keeps its characters."""
    # Given: the existing linear equation input.
    text = "EQTOKEN = 1"
    spec = tmp / "linear.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "eq", "lines": [text], "source": "this document"},
                    {"type": "p", "text": "EQNOTE this document."},
                ]
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "linear.docx"
    before = spec.read_bytes()
    template = sha256(TEMPLATE)
    # When: the real CLI builds it.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the original line is inside Office Math, not a text run, and inputs are unchanged.
    assert result.returncode == 0, result.stderr
    assert spec.read_bytes() == before
    assert sha256(TEMPLATE) == template
    paragraph = _equation_paragraph(output, text)
    assert paragraph.find(f".//{W_NS}t") is None
    omath = paragraph.find(f".//{MATH}oMath")
    assert omath is not None and "".join(omath.itertext()) == text
    assert paragraph.find(f".//{MATH}oMathPara") is not None
    Document(str(output))


def test_latex_source_becomes_native_omml(tmp: Path) -> None:
    """Optional per-line latex becomes structured Office Math and is not visible."""
    # Given: a linear line plus a machine conversion source.
    line = "(4)  SOURCE_LINEAR"
    latex = r"\frac{a}{b}+\sqrt{x}+\sqrt[3]{y}+x_i^{2}+\sum_{i=1} x_i+\sum_{k=1}^{n} k"
    spec = tmp / "native.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "eq", "lines": [line], "latex": [latex], "source": "this document"},
                    {"type": "p", "text": "\uc2dd (4) this document."},
                ]
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "native.docx"
    before = spec.read_bytes()
    # When: the real CLI builds it.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: native structure is present, the number is kept, and the conversion source is absent.
    assert result.returncode == 0, result.stderr
    assert spec.read_bytes() == before
    blob = part(output, "word/document.xml")
    assert latex.encode() not in blob
    assert b"\\frac" not in blob and b"\\sqrt" not in blob and b"\\sum" not in blob
    assert b"SOURCE_LINEAR" not in blob
    paragraph = _equation_paragraph(output, "(4)")
    assert paragraph.find(f".//{W_NS}t") is None
    rads = paragraph.findall(f".//{MATH}rad")
    assert len(rads) == 2
    for radical in rads:
        assert radical.find(f"{MATH}deg") is not None
        expression = radical.find(f"{MATH}e")
        assert expression is not None and any(item.strip() for item in expression.itertext() if item)
    hidden = [radical for radical in rads if radical.find(f"{MATH}radPr/{MATH}degHide") is not None]
    shown = [radical for radical in rads if radical.find(f"{MATH}radPr/{MATH}degHide") is None]
    assert len(hidden) == 1 and hidden[0].find(f"{MATH}radPr/{MATH}degHide").get(f"{MATH}val") == "1"
    assert len(shown) == 1 and "".join(shown[0].find(f"{MATH}deg").itertext()) == "3"
    assert paragraph.find(f".//{MATH}f") is not None
    assert paragraph.find(f".//{MATH}sSub") is not None
    assert paragraph.find(f".//{MATH}sSubSup") is not None or paragraph.find(f".//{MATH}sSup") is not None
    empty_upper = []
    bounded = []
    for operator in paragraph.findall(f".//{MATH}nary"):
        upper = operator.find(f"{MATH}sup")
        assert upper is not None
        if any(upper.itertext()):
            bounded.append(operator)
            assert operator.find(f"{MATH}naryPr/{MATH}supHide") is None
        else:
            empty_upper.append(operator)
            assert operator.find(f"{MATH}naryPr/{MATH}supHide").get(f"{MATH}val") == "1"
    assert len(empty_upper) == 1 and len(bounded) == 1
    Document(str(output))


def test_unsupported_math_does_not_fall_back(tmp: Path) -> None:
    """A conversion source that cannot become complete Office Math writes no document."""
    # Given: conversion sources that leak, break, or omit a radicand.
    samples = (r"\notacommand{1}", r"\frac{1}", r"\sqrt{}")
    for index, latex in enumerate(samples):
        spec = tmp / f"bad-{index}.json"
        spec.write_text(
            json.dumps(
                {
                    "blocks": [
                        {"type": "eq", "lines": ["(8)  kept-linear"], "latex": [latex], "source": "this document"},
                        {"type": "p", "text": "(8) this document."},
                    ]
                }
            ),
            encoding="utf-8",
        )
        output = tmp / f"bad-{index}.docx"
        # When: the real CLI attempts the build.
        result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
        # Then: there is no text equation and the conversion source is not echoed.
        assert result.returncode == 2 and "math_conversion" in result.stderr
        assert not output.exists()
        assert latex not in result.stderr


def _equation_paragraph(path: Path, token: str) -> ET.Element:
    root = ET.fromstring(part(path, "word/document.xml"))
    found = [el for el in root.findall(f".//{W_NS}p") if token in "".join(el.itertext())]
    math = [el for el in found if el.find(f".//{MATH}oMath") is not None]
    assert len(math) == 1
    return math[0]


def test_author_year_does_not_require_bracket_keys(tmp: Path) -> None:
    """Author-year references do not need a fabricated [key] citation."""
    # Given: an in-text author-year fragment and no bracket key.
    text = "The note states the relation (Ada Lovelace, 1843)."
    spec = tmp / "author-year.json"
    spec.write_text(
        json.dumps(
            {
                "citation_style": "author-year",
                "blocks": [{"type": "p", "text": text}],
                "references": [
                    {
                        "key": "Lovelace1843",
                        "cite": "Ada Lovelace, 1843",
                        "text": "Lovelace, A. (1843). A note on an analytical engine. Example Journal, 1, 1-2.",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "author-year.docx"
    # When: the real CLI builds it.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the author-year sentence is preserved and no bracket key was required.
    assert result.returncode == 0, result.stderr
    assert Document(str(output)).paragraphs[0].text == text
    assert b"[Lovelace1843]" not in part(output, "word/document.xml")


def test_author_year_missing_citation_is_rejected(tmp: Path) -> None:
    """An author-year entry that never appears in a paragraph is rejected."""
    # Given: a reference whose cite fragment is absent from the body.
    spec = tmp / "author-year-missing.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [{"type": "p", "text": "No in-text citation appears here."}],
                "references": [
                    {
                        "key": "Lovelace1843",
                        "cite": "Ada Lovelace, 1843",
                        "text": "Lovelace, A. (1843). A note on an analytical engine.",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "author-year-missing.docx"
    # When: the real CLI attempts the build.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the missing citation is rejected and no file is written.
    assert result.returncode == 2 and "missing_reference" in result.stderr
    assert not output.exists()


def test_numbered_style_still_requires_keys(tmp: Path) -> None:
    """Explicit numbered style does not accept an author-year sentence as a citation."""
    # Given: a numbered request whose body has no [key].
    spec = tmp / "numbered.json"
    spec.write_text(
        json.dumps(
            {
                "citation_style": "numbered",
                "blocks": [{"type": "p", "text": "Only an author-year mention (Ada Lovelace, 1843)."}],
                "references": [{"key": "L1", "text": "[L1] Lovelace note."}],
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "numbered.docx"
    # When: the real CLI attempts the build.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: the numbered contract still fails closed.
    assert result.returncode == 2 and "missing_reference" in result.stderr
    assert not output.exists()


def test_callouts_use_source_block_style(tmp: Path) -> None:
    """Learning and summary callouts preserve the source's actual visual style."""
    # Given: two native callout paragraphs in the textbook template.
    spec = tmp / "callout.json"
    spec.write_text(json.dumps({"blocks": [
        {"type": "p", "role": "callout-title", "text": "LEARNING_TITLE"},
        {"type": "p", "role": "callout-body", "text": "LEARNING_BODY"},
    ]}), encoding="utf-8")
    output = tmp / "callout.docx"
    # When: the real CLI builds the manuscript.
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    # Then: both paragraphs use the source callout style, not unstyled prose.
    assert result.returncode == 0, result.stderr
    paragraphs = Document(str(output)).paragraphs
    assert [p.style.style_id for p in paragraphs] == ["BlockText", "BlockText"]
    assert paragraphs[0].runs[0].bold is True
    assert paragraphs[0].paragraph_format.keep_with_next is True
    assert part(output, "word/styles.xml") == part(TEMPLATE, "word/styles.xml")


def test_prose_checker_flags_bad_korean(tmp: Path) -> None:
    """check_prose catches ungrammatical present forms and exact duplicate paragraphs."""
    # Given: one manuscript with defects observed in a reviewer-written chapter, one clean.
    para = "이 문단은 같은 내용을 두 번 싣는지 확인하려고 충분히 길게 쓴 예시 문단이다. 반복되면 독자가 혼란스럽다."
    bad = tmp / "bad.json"
    bad.write_text(json.dumps({"blocks": [
        {"type": "p", "text": "강제 peak 점은 QFlowlevel에서는 빼는다."},
        {"type": "p", "text": "적합은 고정한 sigma를 쓰는다."},
        {"type": "p", "text": para},
        {"type": "p", "text": para},
    ]}, ensure_ascii=False), encoding="utf-8")
    good = tmp / "good.json"
    good.write_text(json.dumps({"blocks": [
        {"type": "p", "text": "강제 peak 점은 QFlowlevel에서 뺀다. 적합은 고정한 sigma를 쓴다. 관측은 먹는다와 같은 받침 동사를 그대로 둔다."},
    ]}, ensure_ascii=False), encoding="utf-8")
    checker = SKILL / "scripts" / "check_prose.py"
    # When: the real checker runs on both.
    bad_run = subprocess.run([sys.executable, str(checker), str(bad)], check=False, capture_output=True, text=True, timeout=60)
    good_run = subprocess.run([sys.executable, str(checker), str(good)], check=False, capture_output=True, text=True, timeout=60)
    # Then: the bad manuscript fails on each defect and the clean one passes.
    assert bad_run.returncode == 1, bad_run.stdout
    assert bad_run.stdout.count("bad_present_form") == 2
    assert "duplicate_paragraph" in bad_run.stdout
    assert good_run.returncode == 0 and "PROSE_OK" in good_run.stdout, good_run.stdout


def test_equation_without_explanation_fails(tmp: Path) -> None:
    """An equation needs a following body paragraph, and a number needs a mention."""
    bare = tmp / "bare.json"
    bare.write_text(
        json.dumps({"blocks": [{"type": "eq", "lines": ["y = 1"], "source": "this document"}]}),
        encoding="utf-8",
    )
    numbered = tmp / "numbered-eq.json"
    numbered.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "eq", "lines": ["(33) y = 1"], "source": "this document"},
                    {"type": "p", "text": "The line is a constant from this document."},
                ]
            }
        ),
        encoding="utf-8",
    )
    bare_out = tmp / "bare.docx"
    numbered_out = tmp / "numbered-eq.docx"
    bare_run = run_cli(["--manuscript", str(bare), "--output", str(bare_out), "--mode", "new"])
    numbered_run = run_cli(["--manuscript", str(numbered), "--output", str(numbered_out), "--mode", "new"])
    assert bare_run.returncode == 2 and "missing_equation_explanation" in bare_run.stderr
    assert numbered_run.returncode == 2 and "missing_equation_explanation" in numbered_run.stderr
    assert not bare_out.exists() and not numbered_out.exists()


def test_equation_source_must_be_in_prose(tmp: Path) -> None:
    """The equation source string must appear in the following body paragraph."""
    spec = tmp / "eq-source.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "eq", "lines": ["y = 1"], "source": "this document"},
                    {"type": "p", "text": "The line is a constant."},
                ]
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "eq-source.docx"
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    assert result.returncode == 2 and "missing_equation_source" in result.stderr
    assert not output.exists()


def test_unknown_author_year_in_source_fails(tmp: Path) -> None:
    """An author-year span in a source must be a reference cite."""
    spec = tmp / "unknown-cite.json"
    spec.write_text(
        json.dumps(
            {
                "citation_style": "author-year",
                "blocks": [
                    {"type": "eq", "lines": ["y = 1"], "source": "Unknown Person, 1999"},
                    {
                        "type": "p",
                        "text": "The line follows Unknown Person, 1999. Known Author, 2020 is listed.",
                    },
                ],
                "references": [
                    {
                        "key": "Known2020",
                        "cite": "Known Author, 2020",
                        "text": "Known Author. (2020). A listed note.",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "unknown-cite.docx"
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    assert result.returncode == 2 and "missing_reference" in result.stderr
    assert not output.exists()


def test_missing_source_file_fails(tmp: Path) -> None:
    """A source path token must exist beside the manuscript."""
    spec = tmp / "missing-file.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "eq", "lines": ["y = 1"], "source": "missing/no-such-file.py"},
                    {"type": "p", "text": "See missing/no-such-file.py."},
                ]
            }
        ),
        encoding="utf-8",
    )
    output = tmp / "missing-file.docx"
    result = run_cli(["--manuscript", str(spec), "--output", str(output), "--mode", "new"])
    assert result.returncode == 2 and "missing_source_file" in result.stderr
    assert not output.exists()


def _table_with_source(tmp: Path, source: str) -> Path:
    spec = tmp / "table.json"
    spec.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "p", "text": f"Rows follow {source}."},
                    {"type": "table", "headers": ["H"], "rows": [["a"]], "source": source},
                ]
            }
        ),
        encoding="utf-8",
    )
    return spec


def test_ledger_missing_entry_fails(tmp: Path) -> None:
    """A ledger must name every table, figure, and numbered equation."""
    spec = _table_with_source(tmp, "this document")
    ledger = tmp / "ledger.json"
    ledger.write_text("[]", encoding="utf-8")
    output = tmp / "ledger-missing.docx"
    result = run_cli(
        ["--manuscript", str(spec), "--ledger", str(ledger), "--output", str(output), "--mode", "new"]
    )
    assert result.returncode == 2 and "missing_ledger_entry" in result.stderr
    assert "table:1" in result.stderr
    assert not output.exists()


def test_example_ledger_emits_native_math(tmp: Path) -> None:
    """The shipped example and ledger build, and the equation is native OMML."""
    output = tmp / "example.docx"
    result = run_cli(
        [
            "--manuscript",
            str(SKILL / "examples" / "manuscript.json"),
            "--ledger",
            str(SKILL / "examples" / "source_ledger.json"),
            "--output",
            str(output),
            "--mode",
            "new",
        ]
    )
    assert result.returncode == 0, result.stderr
    paragraph = _equation_paragraph(output, "y = ax + b")
    assert paragraph.find(f".//{MATH}oMath") is not None
    assert b"layout" not in part(output, "word/document.xml")


def test_unverified_ledger_must_say_so(tmp: Path) -> None:
    """Unverified ledger rows require the in-text source to say so."""
    anchor = tmp / "note.txt"
    anchor.write_text("x", encoding="utf-8")
    admitted = "not verified in this document"
    spec = _table_with_source(tmp, "this document")
    passing = tmp / "pass.json"
    passing.write_text(
        json.dumps(
            {
                "blocks": [
                    {"type": "p", "text": f"Rows follow {admitted}."},
                    {"type": "table", "headers": ["H"], "rows": [["a"]], "source": admitted},
                ]
            }
        ),
        encoding="utf-8",
    )
    missing = tmp / "missing-phrase.json"
    admitted_ledger = tmp / "admitted.json"
    row = {
        "item": "table:1",
        "location": "table",
        "citation": "this document",
        "kind": "convention",
        "checked_against": {"file": str(anchor), "symbol": "note"},
        "status": "unverified",
    }
    missing.write_text(json.dumps([row]), encoding="utf-8")
    admitted_ledger.write_text(json.dumps([row]), encoding="utf-8")
    denied = tmp / "denied.docx"
    allowed = tmp / "allowed.docx"
    denied_run = run_cli(
        ["--manuscript", str(spec), "--ledger", str(missing), "--output", str(denied), "--mode", "new"]
    )
    allowed_run = run_cli(
        [
            "--manuscript",
            str(passing),
            "--ledger",
            str(admitted_ledger),
            "--output",
            str(allowed),
            "--mode",
            "new",
        ]
    )
    assert denied_run.returncode == 2 and "unverified_claimed" in denied_run.stderr
    assert not denied.exists()
    assert allowed_run.returncode == 0, allowed_run.stderr


def test_ledger_literature_needs_reference(tmp: Path) -> None:
    """Literature and adapted ledger rows need a citation that matches a reference."""
    anchor = tmp / "note.txt"
    anchor.write_text("x", encoding="utf-8")
    spec = _table_with_source(tmp, "this document")
    ledger = tmp / "literature.json"
    ledger.write_text(
        json.dumps(
            [
                {
                    "item": "table:1",
                    "location": "table",
                    "citation": "Unknown Person, 1999",
                    "kind": "literature",
                    "checked_against": {"file": str(anchor), "symbol": "note"},
                    "status": "verified",
                }
            ]
        ),
        encoding="utf-8",
    )
    output = tmp / "literature.docx"
    result = run_cli(
        ["--manuscript", str(spec), "--ledger", str(ledger), "--output", str(output), "--mode", "new"]
    )
    assert result.returncode == 2 and "missing_reference" in result.stderr
    assert not output.exists()


def _reject(tmp: Path, source: Path, spec: Path) -> str:
    result = run_cli(
        ["--source", str(source), "--manuscript", str(spec), "--output", str(tmp / f"{spec.stem}.docx"), "--mode", "new"]
    )
    assert result.returncode == 2
    return result.stderr


def main() -> None:
    """Run every check in its own directory."""
    if not TEMPLATE.is_file():
        raise SystemExit(f"missing template: {TEMPLATE}")
    test_help_lists_modes()
    print("PASS help")
    named = (
        ("new", test_new_mode_preserves_style),
        ("reuse", test_new_mode_drops_unrequested_and_keeps_requested),
        ("append", test_append_preserves_package),
        ("plain", test_default_source_without_figures),
        ("reject", test_reject_malformed_missing_and_overwrite),
        ("append-source", test_append_without_source_errors),
        ("windows-path", test_windows_path_is_preserved),
        ("literal-index", test_indexing_is_not_citation),
        ("latex-equation", test_latex_equation_is_rejected),
        ("native-source", test_linear_equation_preserves_source),
        ("native-omml", test_latex_source_becomes_native_omml),
        ("native-reject", test_unsupported_math_does_not_fall_back),
        ("author-year", test_author_year_does_not_require_bracket_keys),
        ("author-year-missing", test_author_year_missing_citation_is_rejected),
        ("numbered-style", test_numbered_style_still_requires_keys),
        ("callout-style", test_callouts_use_source_block_style),
        ("prose-check", test_prose_checker_flags_bad_korean),
        ("eq-explanation", test_equation_without_explanation_fails),
        ("eq-source", test_equation_source_must_be_in_prose),
        ("unknown-cite", test_unknown_author_year_in_source_fails),
        ("missing-file", test_missing_source_file_fails),
        ("ledger-missing", test_ledger_missing_entry_fails),
        ("ledger-example", test_example_ledger_emits_native_math),
        ("ledger-unverified", test_unverified_ledger_must_say_so),
        ("ledger-cite", test_ledger_literature_needs_reference),
    )
    for name, test in named:
        with tempfile.TemporaryDirectory(prefix=f"report-style-{name}-") as tmp:
            test(Path(tmp))
        print(f"PASS {name}")


if __name__ == "__main__":
    main()
